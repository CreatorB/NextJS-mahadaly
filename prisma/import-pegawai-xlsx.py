#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Konversi "DATA PEGAWAI MAHAD ALY.xlsx" -> prisma/data/pegawai.json

Pakai:  python prisma/import-pegawai-xlsx.py [path_xlsx] [path_json]

Skrip ini idempoten: boleh dijalankan ulang setiap kali admin memperbarui
Excel. Hasilnya dipakai oleh `prisma/import-pegawai.ts` (tsx) untuk
menyuntik data ke database.
"""

import json
import os
import re
import sys
from datetime import datetime

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.stderr.write("ERROR: module 'openpyxl' belum terpasang. Jalankan: pip install openpyxl\n")
    sys.exit(1)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_XLSX = os.path.join(
    os.path.dirname(ROOT), "..", "docs", "mahadaly", "DATA PEGAWAI MAHAD ALY.xlsx"
)
DEFAULT_JSON = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "pegawai.json")

HEADER_ROW = 5  # baris judul kolom pada file Google Form

BULAN = {
    "januari": 1, "februari": 2, "maret": 3, "april": 4, "mei": 5, "juni": 6,
    "juli": 7, "agustus": 8, "september": 9, "oktober": 10, "november": 11, "desember": 12,
    "january": 1, "february": 2, "march": 3, "may": 5, "june": 6, "july": 7,
    "august": 8, "september": 9, "october": 10, "november": 11, "december": 12,
}


def bersih(v):
    """Nilai sel -> string rapi; None jadi None."""
    if v is None:
        return None
    s = str(v).strip()
    if not s or s in {"-", "--", "N/A", "n/a"}:
        return None
    return s


def nomor_wa(v):
    """Nomor WA -> bentuk kanonik 08xxxxxxxxx."""
    s = bersih(v)
    if not s:
        return None
    digits = re.sub(r"\D", "", s)
    if not digits:
        return None
    if digits.startswith("00"):
        digits = digits[2:]
    if digits.startswith("62"):
        digits = "0" + digits[2:]
    elif digits.startswith("8") and len(digits) >= 10:
        digits = "0" + digits
    return digits


def parse_lahir(v):
    """'Tempat, Tanggal Lahir' -> (tempat_lahir, 'YYYY-MM-DD' atau None)."""
    s = bersih(v)
    if not s:
        return None, None
    tempat = None
    tanggal = None

    if "," in s:
        bagian = [b.strip() for b in s.split(",", 1)]
        tempat = bagian[0] or None
        sisa = bagian[1]
    else:
        sisa = s

    sisa = sisa.strip()

    # Format 1: DD/MM/YYYY atau DD-MM-YYYY
    m = re.match(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})$", sisa)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            tanggal = datetime(y, mo, d).strftime("%Y-%m-%d")
        except ValueError:
            tanggal = None
        return tempat, tanggal

    # Format 2: "27 Agustus 1986" / "1 Januari 2003"
    m = re.match(r"^(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})$", sisa)
    if m:
        d = int(m.group(1))
        mo = BULAN.get(m.group(2).lower())
        y = int(m.group(3))
        if mo:
            try:
                tanggal = datetime(y, mo, d).strftime("%Y-%m-%d")
            except ValueError:
                tanggal = None
        return tempat, tanggal

    # Format 3: hanya tahun, mis. "Bogor, 1994"
    m = re.search(r"\b(19\d{2}|20\d{2})\b", sisa)
    if m:
        return tempat, None

    return tempat, None


def main():
    xlsx = os.path.abspath(sys.argv[1]) if len(sys.argv) > 1 else os.path.abspath(DEFAULT_XLSX)
    out = os.path.abspath(sys.argv[2]) if len(sys.argv) > 2 else os.path.abspath(DEFAULT_JSON)

    if not os.path.isfile(xlsx):
        sys.stderr.write(f"ERROR: file XLSX tidak ditemukan: {xlsx}\n")
        sys.exit(1)

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < HEADER_ROW:
        sys.stderr.write("ERROR: baris judul kolom tidak ditemukan.\n")
        sys.exit(1)

    header = [bersih(h) for h in rows[HEADER_ROW - 1]]

    def col(name_prefix):
        for i, h in enumerate(header):
            if h and h.lower().startswith(name_prefix.lower()):
                return i
        return None

    idx = {
        "nama": col("Nama (sesuai KTP)"),
        "foto": col("Pas Foto"),
        "jabatan": col("Posisi dan Jabatan"),
        "wa": col("Nomor WA"),
        "email": col("Email"),
        "nik": col("Nomor KTP"),
        "jk": col("Jenis Kelamin"),
        "lahir": col("Tempat, Tanggal Lahir"),
        "pendidikan": col("Pendidikan Terakhir"),
        "ijazah": col("Upload Ijazah"),
        "nikah": col("Status Pernikahan"),
        "alamat": col("Alamat Domisili"),
        "bergabung": col("Tahun Bergabung"),
        "keluarga": col("Anggota Keluarga"),
    }

    if idx["nama"] is None:
        sys.stderr.write("ERROR: kolom 'Nama (sesuai KTP)' tidak ditemukan.\n")
        sys.exit(1)

    def val(row, key):
        i = idx.get(key)
        if i is None or i >= len(row):
            return None
        return row[i]

    hasil = []
    for row in rows[HEADER_ROW:]:
        nama = bersih(val(row, "nama"))
        if not nama:
            continue
        tempat, tgl_lahir = parse_lahir(val(row, "lahir"))
        hasil.append(
            {
                "nama": nama,
                "jabatan": bersih(val(row, "jabatan")),
                "noHp": nomor_wa(val(row, "wa")),
                "email": bersih(val(row, "email")),
                "nik": bersih(val(row, "nik")),
                "jk": bersih(val(row, "jk")),
                "tmpLahir": tempat,
                "tglLahir": tgl_lahir,
                "pendidikan": bersih(val(row, "pendidikan")),
                "statusPernikahan": bersih(val(row, "nikah")),
                "alamat": bersih(val(row, "alamat")),
                "tahunBergabung": bersih(val(row, "bergabung")),
                "kontakKeluarga": bersih(val(row, "keluarga")),
                "fotoUrl": bersih(val(row, "foto")),
                "ijazahUrl": bersih(val(row, "ijazah")),
            }
        )

    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, "w", encoding="utf-8", newline="\n") as f:
        json.dump(hasil, f, ensure_ascii=False, indent=2)
        f.write("\n")

    print(f"SOURCE : {xlsx}")
    print(f"OUTPUT : {out}")
    print(f"PEGAWAI: {len(hasil)}")


if __name__ == "__main__":
    main()
